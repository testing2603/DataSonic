"""Append extraction rows to Excel files in Azure Blob Storage.

Each policy sub-classification maps to one Excel blob under a policy/ folder:
  policy/policy_new_business_extracted.xlsx
  policy/policy_renewal_extracted.xlsx
  policy/policy_cancellation_extracted.xlsx
  policy/policy_mta_extracted.xlsx

Each claims sub-classification maps to an Excel blob under a claims/ folder:
  claims/claims_new_extracted.xlsx
  claims/claims_closure_extracted.xlsx
  claims/claims_update_extracted.xlsx
  claims/claims_reopen_extracted.xlsx

Image metadata (not raw data) is stored in:
  images/document_images.xlsx

Row appends follow the download → modify → re-upload pattern because Excel
files are binary (ZIP) and cannot use Azure Append Blobs.

If AZURE_STORAGE_CONNECTION_STRING is not set, all functions silently no-op
so local-only development works without Azure credentials.
"""

import io
import sys
import threading

import config

from . import extraction_schema as es
from . import claims_extraction_schema as cs

TABLE_TO_BLOB: dict[str, str] = {
    'policy_new_business':  'policy/policy_new_business_extracted.xlsx',
    'policy_renewal':       'policy/policy_renewal_extracted.xlsx',
    'policy_cancellation':  'policy/policy_cancellation_extracted.xlsx',
    'policy_mta':           'policy/policy_mta_extracted.xlsx',
    'claims_new':           'claims/claims_new_extracted.xlsx',
    'claims_closure':       'claims/claims_closure_extracted.xlsx',
    'claims_update':        'claims/claims_update_extracted.xlsx',
    'claims_reopen':        'claims/claims_reopen_extracted.xlsx',
}

IMAGES_BLOB = 'images/document_images.xlsx'
IMAGES_COLUMNS = ['image_id', 'doc_id', 'lob', 'content_type', 'filename', 'blob_url', 'created_at']
IMAGES_FOLDER = 'images/'

_client = None
_blob_locks: dict[str, threading.Lock] = {}
_locks_lock = threading.Lock()


def _lock_for(blob_name: str) -> threading.Lock:
    """Per-blob lock to prevent concurrent download-modify-upload races."""
    with _locks_lock:
        if blob_name not in _blob_locks:
            _blob_locks[blob_name] = threading.Lock()
        return _blob_locks[blob_name]


def _get_container_client():
    global _client
    if _client is not None:
        return _client
    if not config.AZURE_STORAGE_CONNECTION_STRING:
        return None
    from azure.storage.blob import BlobServiceClient
    service = BlobServiceClient.from_connection_string(config.AZURE_STORAGE_CONNECTION_STRING)
    _client = service.get_container_client(config.AZURE_BLOB_CONTAINER)
    return _client


def _get_columns(table_name: str) -> list[str]:
    """Return the column names (excluding auto-increment id) for a table."""
    for _subcat, (tname, columns) in es.SUBCATEGORY_TABLES.items():
        if tname == table_name:
            return [name for name, _, _ in columns if name != 'id']
    for _subcat, (tname, columns) in cs.SUBCATEGORY_TABLES.items():
        if tname == table_name:
            return [name for name, _, _ in columns if name != 'id']
    return []


def _download_workbook(blob_client):
    """Download an existing Excel blob and return an openpyxl Workbook, or None."""
    from openpyxl import load_workbook
    try:
        data = blob_client.download_blob().readall()
        return load_workbook(io.BytesIO(data))
    except Exception:
        return None


def _create_workbook(columns: list[str]):
    """Create a new openpyxl Workbook with a header row."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(columns)
    return wb


def _upload_workbook(blob_client, wb):
    """Serialize a Workbook and upload it (overwrite) as a block blob."""
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    blob_client.upload_blob(buf.getvalue(), overwrite=True)


def ensure_blobs(table_names: list[str] | None = None) -> dict[str, bool]:
    """Pre-create Excel blobs with header rows for the given tables (or all).

    Skips blobs that already exist. Returns {blob_name: created}. Silently
    no-ops when Azure credentials are not configured.
    """
    container = _get_container_client()
    if container is None:
        return {}

    targets = table_names or list(TABLE_TO_BLOB.keys())
    results: dict[str, bool] = {}

    for table_name in targets:
        blob_name = TABLE_TO_BLOB.get(table_name)
        if not blob_name:
            continue
        columns = _get_columns(table_name)
        if not columns:
            continue

        blob_client = container.get_blob_client(blob_name)
        try:
            blob_client.get_blob_properties()
            results[blob_name] = False
        except Exception:
            try:
                wb = _create_workbook(columns)
                _upload_workbook(blob_client, wb)
                results[blob_name] = True
                print(f'Created blob: {blob_name}', file=sys.stderr)
            except Exception as e:
                print(f'Failed to create blob {blob_name}: {e}', file=sys.stderr)
                results[blob_name] = False

    # Ensure images blob exists
    img_blob = container.get_blob_client(IMAGES_BLOB)
    try:
        img_blob.get_blob_properties()
        results[IMAGES_BLOB] = False
    except Exception:
        try:
            wb = _create_workbook(IMAGES_COLUMNS)
            _upload_workbook(img_blob, wb)
            results[IMAGES_BLOB] = True
            print(f'Created blob: {IMAGES_BLOB}', file=sys.stderr)
        except Exception as e:
            print(f'Failed to create blob {IMAGES_BLOB}: {e}', file=sys.stderr)
            results[IMAGES_BLOB] = False

    return results


def append_row(table_name: str, row: dict) -> bool:
    """Append one row to the Excel blob for this table. Returns True on success."""
    container = _get_container_client()
    if container is None:
        return False

    blob_name = TABLE_TO_BLOB.get(table_name)
    if not blob_name:
        return False

    columns = _get_columns(table_name)
    if not columns:
        return False

    blob_client = container.get_blob_client(blob_name)

    with _lock_for(blob_name):
        try:
            wb = _download_workbook(blob_client)
            if wb is None:
                wb = _create_workbook(columns)

            ws = wb.active
            ws.append([row.get(col, '') or '' for col in columns])
            _upload_workbook(blob_client, wb)
            return True
        except Exception as e:
            print(f'Azure Excel append failed for {blob_name}: {e}', file=sys.stderr)
            return False


def upload_image_file(image_id: str, image_data_b64: str,
                      content_type: str) -> str | None:
    """Upload the raw image file to Azure Blob Storage. Returns the blob URL or None."""
    container = _get_container_client()
    if container is None:
        return None

    import base64
    ext = content_type.split('/')[-1] if '/' in content_type else 'png'
    if ext == 'jpeg':
        ext = 'jpg'
    blob_name = f'{IMAGES_FOLDER}{image_id}.{ext}'

    try:
        raw_bytes = base64.b64decode(image_data_b64)
        blob_client = container.get_blob_client(blob_name)
        blob_client.upload_blob(raw_bytes, content_settings=_content_settings(content_type))
        return blob_client.url
    except Exception as e:
        print(f'Azure image upload failed for {blob_name}: {e}', file=sys.stderr)
        return None


def _content_settings(content_type: str):
    from azure.storage.blob import ContentSettings
    return ContentSettings(content_type=content_type)


def append_image_row(image_id: str, doc_id: str, lob: str | None,
                     content_type: str, filename: str | None,
                     blob_url: str | None, created_at: str) -> bool:
    """Append one image metadata row to the images Excel blob."""
    container = _get_container_client()
    if container is None:
        return False

    row_values = {
        'image_id': image_id,
        'doc_id': doc_id,
        'lob': lob or '',
        'content_type': content_type,
        'filename': filename or '',
        'blob_url': blob_url or '',
        'created_at': created_at,
    }

    blob_client = container.get_blob_client(IMAGES_BLOB)

    with _lock_for(IMAGES_BLOB):
        try:
            wb = _download_workbook(blob_client)
            if wb is None:
                wb = _create_workbook(IMAGES_COLUMNS)

            ws = wb.active
            ws.append([row_values.get(col, '') for col in IMAGES_COLUMNS])
            _upload_workbook(blob_client, wb)
            return True
        except Exception as e:
            print(f'Azure Excel append failed for {IMAGES_BLOB}: {e}', file=sys.stderr)
            return False
